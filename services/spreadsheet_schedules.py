"""Bounded spreadsheet parsing without application data or credentials."""
import io,re,zipfile
from datetime import datetime,date,timedelta,time as clock

ALIASES={'date':['дата','date','дата занятия'],'weekday':['день','день недели','weekday'],'number':['пара','номер пары','номер','№ пары','number'],
 'title':['предмет','дисциплина','название','занятие','subject','title'],'start':['начало','время','время начала','start','time'],
 'end':['окончание','конец','время окончания','end'],'teacher':['преподаватель','фио преподавателя','teacher'],'room':['аудитория','кабинет','место','room'],'type':['тип','вид занятия','type']}
WEEKDAYS={'понедельник':0,'пн':0,'вторник':1,'вт':1,'среда':2,'ср':2,'четверг':3,'чт':3,'пятница':4,'пт':4,'суббота':5,'сб':5,'воскресенье':6,'вс':6}

def plain(value):
    if value is None:return ''
    if isinstance(value,(datetime,date,clock)):return value.isoformat()
    return str(value).strip()
def normal(value):return ' '.join(plain(value).lower().replace('ё','е').split())

def workbook_rows(data,filename):
    if len(data)>5*1024*1024:raise ValueError('Файл должен быть не больше 5 МБ.')
    if filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(f.file_size for f in z.infolist())>40*1024*1024 or len(z.infolist())>2000:raise ValueError('Слишком большой файл после распаковки.')
            book=load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
        except ValueError:raise
        except Exception:raise ValueError('Не удалось прочитать XLSX. Сохраните файл в Excel заново.')
        try:
            if len(book.worksheets)>20:raise ValueError('В файле должно быть не больше 20 листов.')
            output={}
            for sheet in book.worksheets:
                if (sheet.max_row or 0)>5000 or (sheet.max_column or 0)>80:raise ValueError('На листе должно быть не больше 5000 строк и 80 столбцов.')
                parsed=[]
                for row in sheet.iter_rows(values_only=True):
                    if len(parsed)>=5000 or len(row)>80:raise ValueError('Слишком большой лист.')
                    parsed.append([plain(cell) for cell in row])
                output[sheet.title]=parsed
            return output
        finally:book.close()
    if filename.lower().endswith('.xls'):
        import xlrd
        try:book=xlrd.open_workbook(file_contents=data,on_demand=True)
        except Exception:raise ValueError('Не удалось прочитать XLS. Сохраните файл как XLSX.')
        try:
            if book.nsheets>20:raise ValueError('Слишком много листов.')
            output={}
            for sheet in book.sheets():
                if sheet.nrows>5000 or sheet.ncols>80:raise ValueError('Слишком большой лист.')
                output[sheet.name]=[[plain(xlrd.xldate.xldate_as_datetime(cell.value,book.datemode)) if cell.ctype==xlrd.XL_CELL_DATE else plain(cell.value) for cell in sheet.row(i)] for i in range(sheet.nrows)]
            return output
        finally:book.release_resources()
    raise ValueError('Загрузите файл .xlsx или .xls.')

def parse_date(value):
    for fmt in ('%Y-%m-%d','%d.%m.%Y','%d/%m/%Y'):
        try:return datetime.strptime(value.split('T')[0],fmt).date()
        except ValueError:pass
    raise ValueError('дата должна содержать день, месяц и год')

def parse_time(value):
    value=value.split('T')[-1].strip().replace('.',':')
    if re.fullmatch(r'\d{1,2}:\d{2}(?::\d{2})?',value):
        hour,minute=map(int,value.split(':')[:2])
        if 0<=hour<=23 and 0<=minute<=59:return f'{hour:02}:{minute:02}'
    raise ValueError('укажите время в формате ЧЧ:ММ')

def preview(data,filename,options):
    sheets=workbook_rows(data,filename);sheet=options.get('sheet') or next(iter(sheets),'');rows=sheets.get(sheet)
    if rows is None:raise ValueError('Лист не найден.')
    header=options.get('header_row')
    if not header:
        header=next((i+1 for i,row in enumerate(rows[:30]) if any(normal(v) in ALIASES['title'] for v in row)),1)
    try:header=int(header)
    except (ValueError,TypeError):raise ValueError('Некорректный номер строки заголовков.')
    if not 1<=header<=min(50,len(rows)):raise ValueError('Строка заголовков не найдена.')
    columns=rows[header-1];mapping={key:next((i for i,value in enumerate(columns) if normal(value) in aliases),None) for key,aliases in ALIASES.items()}
    if isinstance(options.get('mapping'),dict):
        for key,value in options['mapping'].items():
            if key not in ALIASES:continue
            if value is None or value=='':mapping[key]=None
            elif type(value) is int and 0<=value<len(columns):mapping[key]=value
            else:raise ValueError('Некорректное сопоставление столбцов.')
    output=[];errors=[]
    if mapping['title'] is None or mapping['start'] is None or (mapping['date'] is None and mapping['weekday'] is None):errors.append({'row':header,'message':'Выберите столбцы «Предмет», «Начало» и «Дата» или «День недели».'})
    bounds=None
    if mapping['date'] is None and mapping['weekday'] is not None:
        try:
            first=parse_date(options.get('from',''));last=parse_date(options.get('to',''))
            if not 0<=(last-first).days<=366:raise ValueError('Укажите период не длиннее года.')
            bounds=(first,last)
        except ValueError:errors.append({'row':header,'message':'Для дней недели укажите дату начала и окончания периода.'})
    if not errors:
        for index,row in enumerate(rows[header:],start=header+1):
            if not any(row):continue
            def cell(key):
                col=mapping.get(key);return row[col] if col is not None and col<len(row) else ''
            try:
                title=cell('title')
                if not title:raise ValueError('не указан предмет')
                if len(title)>200 or len(cell('teacher'))>200 or len(cell('room'))>120:raise ValueError('слишком длинное название, ФИО или аудитория')
                interval=re.split(r'\s*[-–—]\s*',cell('start'));start=parse_time(interval[0]);end=parse_time(cell('end') or (interval[1] if len(interval)==2 else ''))
                if end<=start:raise ValueError('окончание должно быть позже начала')
                number=cell('number');match=re.fullmatch(r'(\d{1,2})(?:\.0|\s*пара)?',number)
                if number and (not match or not 1<=int(match[1])<=12):raise ValueError('номер пары — от 1 до 12')
                number=int(match[1]) if number else 0
                if mapping['date'] is not None:dates=[parse_date(cell('date'))]
                else:
                    weekday=WEEKDAYS.get(normal(cell('weekday')))
                    if weekday is None:raise ValueError('неизвестный день недели')
                    first,last=bounds;d=first+timedelta(days=(weekday-first.weekday())%7);dates=[]
                    while d<=last:dates.append(d);d+=timedelta(days=7)
                for day in dates:output.append({'date':day.isoformat(),'number':number,'title':title,'start':start,'end':end,'teacher':cell('teacher'),'room':cell('room'),'type':cell('type')[:80] or 'Пара'})
                if len(output)>2000:raise ValueError('в одном импорте допускается не больше 2000 занятий')
            except ValueError as error:errors.append({'row':index,'message':str(error)})
    if not output and not errors:errors.append({'row':header,'message':'Не найдено ни одного занятия.'})
    return {'sheets':list(sheets),'sheet':sheet,'header_row':header,'columns':[{'index':i,'label':value or 'Столбец '+str(i+1)} for i,value in enumerate(columns)],'mapping':mapping,'rows':output[:2000],'errors':errors[:50],'error_count':len(errors)}

