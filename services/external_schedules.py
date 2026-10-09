"""Read public timetable facts from the faculties' own published sources."""
import io
import re
import threading
import time
from collections import defaultdict
from datetime import date, timedelta

import requests
from lxml import html
from openpyxl import load_workbook
from services.schedule_presentation import geography_label


GEO_SOURCE = 'https://docs.google.com/spreadsheets/d/1UlXc5_d5vOx74Pq0O3hCdWD6nRRBuWNZFhBu0ozEsmU/edit'
GEO_EXPORT = GEO_SOURCE.rsplit('/', 1)[0] + '/export?format=xlsx'
HIST_SOURCE = 'http://ccs2.hist.msu.ru/index.php?mnu=143'
HIST_BASE = 'http://ccs2.hist.msu.ru/index.php'
_cache = {}
_cache_lock = threading.Lock()
_days = ('понедельник', 'вторник', 'среда', 'четверг', 'пятница', 'суббота', 'воскресенье')
_geo_week = re.compile(r'(?<![А-Яа-яЁё])(?:(?:по\s+)?(?P<label>верх[а-я]*|ниж[а-я]*)[\s.]*(?:нед[а-я]*)?|(?P<short>в[\s.]*нед[а-я]*))[\s.]*', re.I)
_geo_teacher = re.compile(
    r'(?<![А-Яа-яЁёA-Za-z])(?:ст[.\s]*преп|ст[.\s]*пр|проф(?:ессор)?|доц(?:ент)?|преп|прф|'
    r'акад|внс|снс|нс|пр)[.\s]+([А-ЯЁA-Z][А-ЯЁа-яёA-Za-z-]+(?:\s+[А-ЯЁA-Z]\.\s*[А-ЯЁA-Z]\.)?)',
    re.I,
)
_geo_groups = re.compile(r'\b([12])\s*п/[гп]\s*(\d{2,4})\b', re.I)
_geo_spaced_titles = {
    'ГИДРОЛОГИЯ': 'Гидрология', 'БИОЛОГИЯ': 'Биология',
    'ТОПОГРАФИЯ': 'Топография', 'МАТЕМАТИКА': 'Математика',
    'ИСТОРИЯРОССИИ': 'История России', 'СПЕЦПОДГОТОВКА': 'Спецподготовка',
    'ФИЛОСОФИЯ': 'Философия',
}


def _get(url, params=None, limit=5_000_000, session=None):
    response = (session or requests).get(url, params=params, timeout=25,
                                         headers={'User-Agent': 'FGU-Diary/1.0 (public timetable attribution)'})
    response.raise_for_status()
    if len(response.content) > limit:
        raise ValueError('Расписание превышает допустимый размер')
    return response.content


def _cached(key, ttl, build):
    with _cache_lock:
        value = _cache.get(key)
        if value and time.monotonic() - value[0] < ttl:
            return value[1]
    result = build()
    with _cache_lock:
        _cache[key] = (time.monotonic(), result)
    return result


def _text(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def _weekday(value):
    normal = re.sub(r'[^а-яё]', '', _text(value).casefold())
    return next((index for index, name in enumerate(_days) if normal == name), None)


def _times(value):
    match = re.search(r'(\d{1,2})\s*[-.:]\s*(\d{2})\s*[-–—]\s*(\d{1,2})\s*[-.:]\s*(\d{2})', _text(value))
    if not match:
        return None
    a, b, c, d = map(int, match.groups())
    if a > 23 or c > 23 or b > 59 or d > 59:
        return None
    return f'{a:02}:{b:02}', f'{c:02}:{d:02}'


def _period(number, title, start, end, kind='Пара', room='', teacher=''):
    return {'number': number, 'periods': [{
        'disciplineFullName': title[:300], 'timeStart': start, 'timeEnd': end,
        'teachersNameFull': teacher[:200], 'classroom': room[:120], 'groups': '', 'typeStr': kind[:80],
    }]}


def _geo_is_metadata(value):
    remaining = _geo_teacher.sub('', value)
    remaining = re.sub(r'\b\d{1,4}\b|[/,.;()\s-]', '', remaining)
    return not remaining or remaining.casefold() in ('час', 'часа', 'часов')


def _geo_variants(raw):
    markers = [match for match in _geo_week.finditer(raw)
               if raw[:match.start()].count('(') == raw[:match.start()].count(')')]
    if not markers:
        return [(None, raw)]
    parity = lambda match: 1 if (match.group('label') or '').casefold().startswith('ниж') else 0
    prefix = raw[:markers[0].start()].strip(' ,;/')
    if len(markers) == 1:
        suffix = raw[markers[0].end():].strip(' ,;/')
        if prefix and _geo_is_metadata(suffix):
            preceding_room = re.search(r'\b\d{2,4}\b\s*$', prefix)
            if preceding_room and suffix:
                base = prefix[:preceding_room.start()].strip(' ,;/')
                return [(1 - parity(markers[0]), prefix),
                        (parity(markers[0]), _text(base + ' ' + suffix))]
        return [(parity(markers[0]), _text(prefix + ' ' + suffix))]
    segments = [raw[marker.end():markers[index + 1].start() if index + 1 < len(markers) else len(raw)].strip(' ,;/')
                for index, marker in enumerate(markers)]
    if prefix and all(_geo_is_metadata(segment) for segment in segments):
        return [(parity(marker), _text(prefix + ' ' + segment)) for marker, segment in zip(markers, segments)]
    if prefix:
        pieces = [prefix] + [segment for segment in segments if segment]
        return [(parity(marker), piece) for marker, piece in zip(markers, pieces)]
    return [(parity(marker), segment) for marker, segment in zip(markers, segments) if segment]


def _geo_title(value):
    value = _text(value).strip(' .,;/-')
    if re.fullmatch(r'(?:[А-ЯЁ]\s+){3,}[А-ЯЁ]', value):
        return _geo_spaced_titles.get(value.replace(' ', ''), value.replace(' ', '').capitalize())
    if value.upper() == value and re.search('[А-ЯЁ]', value):
        if len(value) <= 4:
            return value
        value = re.sub(r'\s*-\s*', '-', value.lower().capitalize())
        return re.sub(r'\b(россии|россия|мгу|гис|фгм|рф|сша)\b',
                      lambda match: {'россии': 'России', 'россия': 'Россия'}.get(
                          match.group(), match.group().upper()), value)
    return value


def _geo_periods(raw, start, end):
    # Meeting links are timetable metadata, never part of a subject or room.
    raw = re.sub(r'(?:https?://|www\.)[^\s<>]+', ' ', raw, flags=re.I)
    kind = 'Лекция' if re.search(r'\(\s*л\s*\)', raw, re.I) else ('Семинар' if re.search(r'\(\s*с\s*\)', raw, re.I) else 'Пара')
    teachers = []
    def remove_teacher(match):
        name = match.group(1)
        if name.startswith('C') and len(name) > 1 and 'А' <= name[1] <= 'я':
            name = 'С' + name[1:]
        teachers.append(name.title() if name.upper() == name else name)
        return ' '
    clean = _geo_teacher.sub(remove_teacher, raw)
    clean = _geo_week.sub(' ', clean)
    clean = re.sub(r'\bс\s+\d{1,2}[.:]\d{2}\b|\b\d\s*час(?:а|ов)?\b', ' ', clean, flags=re.I)
    clean = re.sub(r'\(\s*[лс]\s*\)', ' ', clean, flags=re.I)
    teacher = ', '.join(dict.fromkeys(teachers))[:200]
    def build(title, room='', group=''):
        label, room = geography_label(title, room, source=raw)
        return {'disciplineFullName': _geo_title(title)[:300], 'timeStart': start, 'timeEnd': end,
                '_diary_title': label, 'teachersNameFull': teacher, 'classroom': room[:120], 'groups': group, 'typeStr': kind}
    if re.match(r'^\s*Иностранный\s+язык\b', clean, re.I):
        rooms = list(dict.fromkeys(re.findall(r'(?<!\d)\d{4}(?!\d)', clean)))
        return [build('Иностранный язык', ', '.join(rooms[:12]))]
    groups = list(_geo_groups.finditer(clean))
    if groups:
        periods = []
        previous = 0
        title = ''
        for match in groups:
            candidate = _geo_title(clean[previous:match.start()])
            if candidate:
                title = candidate
            if title:
                periods.append(build(title, match.group(2), match.group(1) + ' п/г'))
            previous = match.end()
        if periods:
            return periods
    rooms = list(dict.fromkeys(re.findall(r'(?<!\d)\d{2,4}(?!\d)', clean)))
    title = re.sub(r'(?<!\d)\d{2,4}(?!\d)', ' ', clean)
    title = _geo_title(title)
    return [build(title, ', '.join(rooms[:6]))] if title else []


def _lesson_days(lessons):
    result = defaultdict(list)
    for key, day, lesson in lessons:
        result[key, day].append(lesson)
    grouped = defaultdict(list)
    for (key, day), rows in result.items():
        rows.sort(key=lambda row: row['periods'][0]['timeStart'])
        grouped[key].append({'date': day, 'lessons': rows})
    for rows in grouped.values():
        rows.sort(key=lambda row: row['date'])
    return dict(grouped)


def parse_geo_workbook(data):
    book = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    catalog = []
    entries = []
    for sheet in book:
        title = _text(sheet.cell(1, 1).value or sheet.cell(1, 4).value)
        if 'ГЕОГРАФИЧЕСКИЙ ФАКУЛЬТЕТ' not in title.upper():
            continue
        match = re.search(r'(\d)\s*КУРС', title.upper())
        if not match:
            continue
        course = int(match[1])
        level = 'Магистратура' if 'МАГИСТРАТУР' in title.upper() else 'Бакалавриат'
        year_match = re.search(r'20\d{2}\s*[-–]\s*20\d{2}', title)
        term_year = int(year_match.group()[:4]) if year_match else date.today().year
        spring = 'ВЕСЕНН' in title.upper()
        first = date(term_year + 1, 2, 1) if spring else date(term_year, 9, 1)
        last = date(term_year + 1, 6, 30) if spring else date(term_year, 12, 31)
        first_monday = first - timedelta(days=first.weekday())
        header = 2 if course == 1 and level == 'Бакалавриат' else (3 if course == 2 and level == 'Бакалавриат' else 2)
        end_col = next((col for col in range(4, min(sheet.max_column, 40)) if _text(sheet.cell(header, col).value).casefold() == 'пара'), sheet.max_column)
        names = {}
        used = defaultdict(int)
        current_program = ''
        for col in range(4, end_col):
            if header == 3:
                current_program = _text(sheet.cell(2, col).value) or current_program
            cell = sheet.cell(header, col)
            name = _text(cell.value)
            if not name and header == 3:
                name = _text(sheet.cell(2, col).value)
            if not name or name.casefold() == 'пара':
                continue
            if re.fullmatch(r'\d+\.0', name):
                name = name[:-2]
            used[name] += 1
            display = f'{name} ({used[name]})' if used[name] > 1 else name
            key = f'geo:{sheet.title}:{col}'
            names[col] = key
            catalog.append({'id': key, 'faculty': 'geo', 'source_id': key, 'name': display,
                            'level': level, 'course': course, 'program': current_program if header == 3 else '',
                            'source': 'geo'})
        merges = {}
        for area in sheet.merged_cells.ranges:
            if area.max_col >= 4 and area.min_col < end_col and area.min_row > header:
                for row in range(area.min_row, area.max_row + 1):
                    for col in range(max(4, area.min_col), min(end_col - 1, area.max_col) + 1):
                        merges[row, col] = (area.min_row, area.min_col)
        current_day = 0
        current_time = None
        for row in range(header + 1, min(sheet.max_row, 300) + 1):
            day = _weekday(sheet.cell(row, 1).value)
            if day is not None:
                current_day = day
            interval = _times(sheet.cell(row, 2).value)
            if interval:
                current_time = interval
            number_text = _text(sheet.cell(row, 3).value).casefold()
            number_match = re.match(r'(\d{1,2})(?:\.0|\s*верх\w*|\s*ниж\w*)?$', number_text)
            if not number_match or not current_time:
                continue
            number = int(number_match[1])
            if not 1 <= number <= 12:
                continue
            row_parity = 0 if 'верх' in number_text else (1 if 'ниж' in number_text else None)
            for col, key in names.items():
                origin = merges.get((row, col), (row, col))
                raw = _text(sheet.cell(*origin).value)
                if not raw or raw.casefold() in ('день самостоятельной работы', 'научно-исследовательская работа'):
                    continue
                variants = []
                for marked_parity, fragment in _geo_variants(raw):
                    periods = _geo_periods(fragment, *current_time)
                    if periods and (marked_parity is None or row_parity is None or marked_parity == row_parity):
                        variants.append((marked_parity if marked_parity is not None else row_parity, periods))
                current = first + timedelta(days=(current_day - first.weekday()) % 7)
                while current <= last:
                    parity = (current - first_monday).days // 7 % 2
                    for cell_parity, periods in variants:
                        if cell_parity is None or cell_parity == parity:
                            entries.append((key, current.isoformat(), {'number': number, 'periods': periods}))
                    current += timedelta(days=7)
    return catalog, _lesson_days(entries)


def geo_data():
    return _cached('geo', 1800, lambda: parse_geo_workbook(_get(GEO_EXPORT)))


def _hist_tree(data):
    return html.fromstring(data.decode('cp1251', errors='replace'))


def parse_hist_groups(data, level, course, faculty_code, year_id, program='', program_id=''):
    page = _hist_tree(data)
    tables = page.xpath('//table[@id="TblSELVZ"]')
    if not tables:
        raise ValueError('Сайт истфака изменил список групп')
    rows = []
    for anchor in tables[0].xpath('.//a[@href]'):
        match = re.fullmatch(r'\?gr=(\d+)', anchor.get('href', ''))
        name = _text(anchor.text_content())
        if not match or not name or name == '.':
            continue
        source_id = f'{faculty_code}:{year_id}:{program_id}:{match[1]}'
        rows.append({'id': 'hist:' + match[1], 'faculty': 'hist', 'source_id': source_id,
                     'name': name, 'level': level, 'course': course, 'program': program, 'source': 'hist'})
    return rows


def _hist_catalog():
    catalog = {}
    levels = ((1, 'Бакалавриат', 4), (2, 'Магистратура', 2),
              (3, 'Специалитет', 6), (4, 'Интегрированная магистратура', 6))
    academic_year = date.today().year if date.today().month >= 9 else date.today().year - 1
    with requests.Session() as session:
        _get(HIST_BASE, {'mnu': 143}, session=session)
        for faculty_code, level, max_course in levels:
            base = _hist_tree(_get(HIST_BASE, {'f': faculty_code}, session=session))
            programs = [(re.fullmatch(r'\?sp=(\d+)', anchor.get('href', '')), _text(anchor.text_content()).strip('[]'))
                        for anchor in base.xpath('//table[@id="TblSELVZ"]//a[@href]')]
            programs = [(match[1], name) for match, name in programs if match and name] or [('', '')]
            for program_id, program in programs:
                page = _hist_tree(_get(HIST_BASE, {'sp': program_id}, session=session)) if program_id else base
                options = page.xpath('//table[@id="TblSELVZ"]//select[@name="yr"]/option')
                for option in options:
                    year_match = re.search(r'20\d{2}', option.text_content())
                    year_id = option.get('value', '')
                    if not year_match or not year_id.isdigit():
                        continue
                    course = academic_year - int(year_match.group()) + 1
                    if not 1 <= course <= max_course:
                        continue
                    listing = _get(HIST_BASE, {'yr': year_id}, session=session)
                    for row in parse_hist_groups(listing, level, course, faculty_code, year_id, program, program_id):
                        catalog[row['id']] = row
    if not catalog:
        raise ValueError('Сайт истфака не вернул ни одной группы')
    return list(catalog.values())


def hist_groups():
    return _cached('hist-groups', 21600, _hist_catalog)


def parse_hist_schedule(data, expected_group=None):
    page = _hist_tree(data)
    if expected_group:
        selected = page.xpath('//table[@id="TblSELVZ"]//a[@href="?gr='+str(expected_group)+'"]/b')
        if not selected:
            raise ValueError('Сайт истфака переключил выбранную группу')
    tables = page.xpath('//table[@bordercolor="#FF0000"]')
    result = []
    slot_times = {}
    for table in tables:
        for row in table.xpath('./tr|./tbody/tr')[1:]:
            cells = row.xpath('./td')
            if cells and _text(cells[0].text_content()).isdigit():
                interval = _times(cells[0].get('title'))
                if interval:
                    slot_times[int(_text(cells[0].text_content()))] = interval
    for table in tables:
        rows = table.xpath('./tr|./tbody/tr')
        if len(rows) < 2:
            continue
        date_match = re.search(r'\b(\d{2})\.(\d{2})\.(20\d{2})\b', _text(rows[0].text_content()))
        if not date_match:
            continue
        day = date(int(date_match[3]), int(date_match[2]), int(date_match[1])).isoformat()
        lessons = []
        for number, row in enumerate(rows[1:], 1):
            cells = row.xpath('./td')
            if not cells:
                continue
            interval = slot_times.get(number)
            if not interval:
                continue
            for item in cells[-1].xpath('.//div[@id="LESS"]'):
                description = _text(item.get('title'))
                visible = _text(item.text_content())
                subject = re.search(r"['‘](.+?)['’]", description)
                title = subject[1] if subject else visible
                kind = description.split("'")[0].strip() or 'Пара'
                bracket = re.search(r'\[\s*([^]]+)\s*\]', visible)
                room = _text(visible.split('[', 1)[0].replace(_text(item.xpath('string(.//b[1])')), '', 1))
                teacher = _text(visible[bracket.end():]) if bracket else ''
                lessons.append(_period(number, title, *interval, kind=kind, room=room, teacher=teacher))
        if lessons:
            result.append({'date': day, 'lessons': lessons})
    return result


def _load_hist_schedule(source_id, today):
    parts = source_id.split(':')
    if len(parts) == 3:
        faculty_code, year_id, group = parts
        program_id = ''
    else:
        faculty_code, year_id, program_id, group = parts
    now = today
    months = [(now.year + (now.month + offset - 1) // 12, (now.month + offset - 1) % 12 + 1)
              for offset in range(3)]
    result = []
    with requests.Session() as session:
        _get(HIST_BASE, {'mnu': 143}, session=session)
        _get(HIST_BASE, {'f': faculty_code}, session=session)
        if program_id:
            _get(HIST_BASE, {'sp': program_id}, session=session)
        _get(HIST_BASE, {'yr': year_id}, session=session)
        _get(HIST_BASE, {'gr': group}, session=session)
        for year, month in months:
            data = _get(HIST_BASE, {'pMns': f'{month}.{year}'}, session=session)
            result.extend(parse_hist_schedule(data, expected_group=group))
    return result


def hist_schedule(source_id, today=None):
    now = today or date.today()
    return _cached(f'hist-schedule:{source_id}:{now.year}-{now.month}', 1800,
                   lambda: _load_hist_schedule(source_id, now))
